from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


def _add_section(document: Document, title_text: str, requirements: list[dict]) -> None:
    heading = document.add_paragraph()
    run = heading.add_run(title_text)
    run.bold = True
    run.font.size = Pt(14)

    columns = [
        "ID",
        "Requirement",
        "Priority",
        "Rationale",
        "Acceptance Criteria",
    ]

    table = document.add_table(rows=1, cols=len(columns))
    table.style = "Light List Accent 1"
    hdr_cells = table.rows[0].cells
    for idx, name in enumerate(columns):
        hdr_cells[idx].text = name

    for req in requirements:
        row = table.add_row().cells
        row[0].text = req.get("ID", "")
        row[1].text = req.get("Requirement", "")
        row[2].text = req.get("Priority", "")
        row[3].text = req.get("Rationale", "")
        row[4].text = req.get("Acceptance Criteria", "")

    document.add_paragraph("")


def create_system_requirements_docx(output_path: str) -> None:
    document = Document()

    # Title
    title = document.add_paragraph()
    run = title.add_run("System Requirements Specification (SRS)")
    run.font.size = Pt(20)
    run.bold = True
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Metadata subtitle
    subtitle = document.add_paragraph()
    subtitle_run = subtitle.add_run(
        "Project: <Your Project Title>\nVersion: 1.0\nDate: <YYYY-MM-DD>\nAuthors: <Your Name>"
    )
    subtitle_run.font.size = Pt(11)

    document.add_paragraph("")

    # Instructions
    instr = document.add_paragraph()
    instr_run = instr.add_run(
        "Instructions: Replace angle-bracketed placeholders and adjust sample requirements as needed. "
        "Use the table below to maintain a concise, thesis-ready requirements list."
    )
    instr_run.italic = True

    document.add_paragraph("")

    # Functional Requirements
    functional_reqs = [
        {
            "ID": "FR-01",
            "Requirement": "Students shall authenticate using university credentials.",
            "Priority": "Must",
            "Rationale": "Access control to protect data.",
            "Acceptance Criteria": "Valid credentials log in within 2 seconds; invalid show clear error.",
        },
        {
            "ID": "FR-02",
            "Requirement": "Display a personalized dashboard with enrolled courses and progress.",
            "Priority": "Must",
            "Rationale": "Primary student use case.",
            "Acceptance Criteria": "Post-login, dashboard lists courses with progress bars and timestamps.",
        },
        {
            "ID": "FR-03",
            "Requirement": "Admins can add, edit, and remove courses, departments, and users.",
            "Priority": "Must",
            "Rationale": "Ongoing content administration.",
            "Acceptance Criteria": "CRUD persists to DB and updates UI without errors.",
        },
    ]
    _add_section(document, "Functional Requirements", functional_reqs)

    # Non-Functional Requirements
    nonfunctional_reqs = [
        {
            "ID": "NFR-01",
            "Requirement": "Dashboard loads within 3 seconds for 95% of nominal-load requests.",
            "Priority": "Must",
            "Rationale": "Usability and performance.",
            "Acceptance Criteria": "Load tests show p95 < 3s for dashboard endpoint.",
        },
        {
            "ID": "NFR-02",
            "Requirement": "Encrypt all data in transit with TLS 1.2+.",
            "Priority": "Must",
            "Rationale": "Security and compliance.",
            "Acceptance Criteria": "Security scan: HTTPS only, no mixed content, HSTS (if applicable).",
        },
        {
            "ID": "NFR-03",
            "Requirement": "Keyboard-accessible UI with alt text for images (WCAG 2.1 AA checks).",
            "Priority": "Should",
            "Rationale": "Accessibility.",
            "Acceptance Criteria": "Audit confirms keyboard navigation and alt text coverage.",
        },
        {
            "ID": "NFR-04",
            "Requirement": "Monthly uptime of 99.5% excluding planned maintenance.",
            "Priority": "Should",
            "Rationale": "Reliability expectations.",
            "Acceptance Criteria": "Monitoring shows uptime >= 99.5% in a calendar month.",
        },
    ]
    _add_section(document, "Non-Functional Requirements", nonfunctional_reqs)

    # Technical Requirements
    technical_reqs = [
        {
            "ID": "TECH-01",
            "Requirement": "Use Python 3.10+ and Flask for the web backend.",
            "Priority": "Must",
            "Rationale": "Project stack alignment.",
            "Acceptance Criteria": "Application starts and passes smoke tests on Python 3.10+.",
        },
        {
            "ID": "TECH-02",
            "Requirement": "Use SQLite for the thesis demo environment with SQLAlchemy ORM.",
            "Priority": "Must",
            "Rationale": "Simplicity and portability.",
            "Acceptance Criteria": "DB file created in instance directory; CRUD operations succeed.",
        },
        {
            "ID": "TECH-03",
            "Requirement": "Implement role-based access control for Student, Faculty, Admin.",
            "Priority": "Must",
            "Rationale": "Security and separation of duties.",
            "Acceptance Criteria": "Unauthorized access returns 403; authorized flows succeed.",
        },
    ]
    _add_section(document, "Technical Requirements", technical_reqs)

    # Hardware Requirements
    hardware_reqs = [
        {
            "ID": "HW-01",
            "Requirement": "Server (demo): x86_64 CPU, 2 cores, 4 GB RAM, 10 GB storage.",
            "Priority": "Should",
            "Rationale": "Capacity for demo workloads.",
            "Acceptance Criteria": "System deploys and handles nominal demo usage without resource exhaustion.",
        },
        {
            "ID": "HW-02",
            "Requirement": "Client: modern laptop/desktop capable of running a current browser.",
            "Priority": "Must",
            "Rationale": "Ensure smooth UI performance.",
            "Acceptance Criteria": "UI interaction is responsive on typical university hardware.",
        },
    ]
    _add_section(document, "Hardware Requirements", hardware_reqs)

    # Software Requirements
    software_reqs = [
        {
            "ID": "SW-01",
            "Requirement": "Operating system support: Windows 10+, macOS 12+, Ubuntu 22.04+.",
            "Priority": "Could",
            "Rationale": "Portability for evaluation environments.",
            "Acceptance Criteria": "Install and smoke tests pass on the three OS families.",
        },
        {
            "ID": "SW-02",
            "Requirement": "Browsers: latest Chromium, Firefox, or Safari within last 12 months.",
            "Priority": "Must",
            "Rationale": "Security and modern web features.",
            "Acceptance Criteria": "Manual checks confirm core flows work on stated versions.",
        },
        {
            "ID": "SW-03",
            "Requirement": "Dependencies: python-docx, SQLAlchemy, Flask and listed requirements.txt packages.",
            "Priority": "Must",
            "Rationale": "App functionality and document generation.",
            "Acceptance Criteria": "pip install -r requirements.txt completes; app starts successfully.",
        },
    ]
    _add_section(document, "Software Requirements", software_reqs)

    notes = document.add_paragraph()
    notes_run = notes.add_run(
        "Notes: Add or refine rows per your thesis context. Keep requirements atomic and testable."
    )
    notes_run.italic = True

    document.save(output_path)


if __name__ == "__main__":
    output_file = "System_Requirements.docx"
    create_system_requirements_docx(output_file)



